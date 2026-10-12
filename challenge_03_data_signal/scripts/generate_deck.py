"""Generate slides/presentation_deck.pptx for Challenge 03."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = Path(__file__).resolve().parent.parent
FIG = BASE / "reports" / "figures"
OUT = BASE / "slides" / "presentation_deck.pptx"

NAVY = RGBColor(0x1b, 0x4f, 0x72)
BLUE = RGBColor(0x28, 0x74, 0xa6)
AMBER = RGBColor(0xe6, 0x7e, 0x22)
RED = RGBColor(0xc0, 0x39, 0x2b)
GREEN = RGBColor(0x1e, 0x84, 0x49)
INK = RGBColor(0x1a, 0x23, 0x32)
SUB = RGBColor(0x5b, 0x64, 0x72)
WHITE = RGBColor(0xff, 0xff, 0xff)
PALE = RGBColor(0xd6, 0xea, 0xf8)
LIGHT = RGBColor(0xee, 0xf4, 0xf9)
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, content, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    if isinstance(content, str):
        content = [content]
    for i, para in enumerate(content if len(content) else [""]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(para, tuple):
            txt, sty = para
        else:
            txt, sty = para, {}
        run = p.add_run(); run.text = txt
        run.font.size = Pt(sty.get("size", size))
        run.font.bold = sty.get("bold", bold)
        run.font.color.rgb = sty.get("color", color)
        run.font.name = FONT
    return tb


def lines(s, x, y, w, h, items, size=14, gap=0.55):
    """Render (label, rest) pairs as bold-label + normal-rest lines."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(4)
        r1 = p.add_run(); r1.text = lead
        r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = NAVY; r1.font.name = FONT
        if rest:
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size - 1); r2.font.color.rgb = INK; r2.font.name = FONT
    return tb


def title_bar(s, title, subtitle=""):
    rect(s, 0, 0, prs.slide_width, Inches(1.0), NAVY)
    text(s, Inches(0.55), Inches(0.18), Inches(11), Inches(0.6), title,
         size=26, color=WHITE, bold=True)
    if subtitle:
        text(s, Inches(0.57), Inches(0.72), Inches(12), Inches(0.4), subtitle,
             size=13, color=PALE)


def bullets(s, x, y, w, h, items, size=15):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run(); r.text = "• " + lead
        r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = FONT
        if rest:
            r2 = p.add_run(); r2.text = rest
            r2.font.size = Pt(size - 1); r2.font.color.rgb = INK; r2.font.name = FONT
    return tb


def pic(s, path, x, y, w):
    s.shapes.add_picture(str(path), x, y, width=w)


# ============================================================================
# 1 — Title
# ============================================================================
s = slide()
rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
text(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.0),
     "挑战 03 · 数据 → 信号", size=44, color=WHITE, bold=True)
text(s, Inches(0.9), Inches(3.3), Inches(11.5), Inches(0.9),
     "把干净的数据转化为启发性的预测信号 —— 季节性 · 周期 · 节假日",
     size=20, color=PALE)
text(s, Inches(0.9), Inches(4.5), Inches(11), Inches(0.5),
     "《杂货零售——数据驱动》第三里程碑：可靠性与远见的连接", size=14, color=PALE)
for i, (v, l) in enumerate([("~365d / 7d", "无监督发现的主导周期"),
                            ("+12.5%", "节假日销售抬升"),
                            ("10.5%", "总序列 MAPE"),
                            ("78%→30%", "Bakery 品类 MAPE")]):
    x = Inches(0.9 + i * 3.05)
    text(s, x, Inches(5.45), Inches(2.9), Inches(0.7), v, size=30, color=WHITE, bold=True)
    text(s, x, Inches(6.15), Inches(2.9), Inches(0.4), l, size=13, color=PALE)

# 2 — Objectives
s = slide()
title_bar(s, "目标：让数据「有启发性」", "承接挑战 02 的诚信，进入挑战 03 的信号")
bullets(s, Inches(0.55), Inches(1.5), Inches(6.5), Inches(5.2), [
    ("探索季节性", " —— 周度与年度的稳定节奏"),
    ("识别需求周期", " —— 无监督周期图自动发现"),
    ("量化外部因子", " —— 节假日与价格"),
    ("构建可解释预测流程", " —— 透明、可复现、业务意识"),
], size=16)
rect(s, Inches(7.4), Inches(1.5), Inches(5.4), Inches(4.9), LIGHT)
text(s, Inches(7.7), Inches(1.75), Inches(4.8), Inches(0.5), "数据范围", size=18, color=NAVY, bold=True)
lines(s, Inches(7.7), Inches(2.35), Inches(4.9), Inches(4.0), [
    ("5 年", "  2018–2022（1826 天，含闰日）"),
    ("3 门店", "  西雅图 / 波特兰 / 博伊西"),
    ("4 品类", "  Beverages / Bakery / Meat / Produce"),
    ("", ""),
    ("21,912 条", "  完整日网格，无缺口"),
])

# 3 — Unsupervised seasonality
s = slide()
title_bar(s, "无监督发现：周 + 年双周期", "FFT 周期图，零日历先验")
pic(s, FIG / "fig2_periodogram.png", Inches(0.55), Inches(1.5), w=Inches(8.1))
rect(s, Inches(8.9), Inches(1.55), Inches(3.9), Inches(4.8), LIGHT)
text(s, Inches(9.15), Inches(1.8), Inches(3.5), Inches(0.6), "季节强度（STL）", size=17, color=NAVY, bold=True)
lines(s, Inches(9.15), Inches(2.45), Inches(3.6), Inches(3.2), [
    ("Produce 0.33", "  最敏感"),
    ("Beverages 0.31", ""),
    ("Bakery 0.28", ""),
    ("Meat 0.24", "  最刚性"),
])
text(s, Inches(9.15), Inches(5.2), Inches(3.6), Inches(0.6), "总序列 0.38", size=16, color=AMBER, bold=True)

# 4 — Seasonal & external
s = slide()
title_bar(s, "季节模式与外部因子", "周/年节奏 + 节假日脉冲 + 价格信号")
pic(s, FIG / "fig4_annual_seasonality.png", Inches(0.55), Inches(1.5), w=Inches(6.3))
pic(s, FIG / "fig7_holiday_uplift.png", Inches(7.0), Inches(1.5), w=Inches(3.0))
rect(s, Inches(10.2), Inches(1.55), Inches(2.55), Inches(4.6), LIGHT)
text(s, Inches(10.4), Inches(1.75), Inches(2.2), Inches(0.6), "价格弹性", size=16, color=NAVY, bold=True)
text(s, Inches(10.4), Inches(2.35), Inches(2.3), Inches(3.2), "≈ 0", size=30, color=RED, bold=True)
text(s, Inches(10.4), Inches(3.2), Inches(2.3), Inches(1.6), "|r| < 0.02\n需求由日历驱动\n而非价格",
     size=12, color=SUB)
lines(s, Inches(0.55), Inches(5.35), Inches(9.5), Inches(1.6), [
    ("Q4 节假日峰", "  Bakery 最陡峭；节假日销售均值 +12.5%"),
    ("无监督异常检测", "  18 个异常点中 12 个命中节假日（精确率 66.7%）"),
])

# 5 — Forecast
s = slide()
title_bar(s, "预测：SARIMA → SARIMAX", "训练 2018–2021，全盲测 2022（365 天）")
rect(s, Inches(0.55), Inches(1.5), Inches(3.6), Inches(4.6), LIGHT)
text(s, Inches(0.75), Inches(1.65), Inches(3.0), Inches(0.5), "总序列 MAPE", size=15, color=NAVY, bold=True)
rows = [("persistence", "24.7%", False),
        ("seasonal-naive", "20.5%", False),
        ("SARIMA（仅周）", "15.1%", False),
        ("SARIMAX（+假日/年度）", "10.5%", True)]
yy = 2.25
for name, val, hl in rows:
    text(s, Inches(0.75), Inches(yy), Inches(2.2), Inches(0.5), name,
         size=13, color=(RED if hl else INK), bold=hl)
    text(s, Inches(2.95), Inches(yy), Inches(1.1), Inches(0.5), val,
         size=13, color=(RED if hl else INK), bold=hl, align=PP_ALIGN.RIGHT)
    yy += 0.72
pic(s, FIG / "fig10_forecast_total.png", Inches(4.4), Inches(1.5), w=Inches(8.3))

# 6 — Hierarchy
s = slide()
title_bar(s, "分层结果：外部因子全面生效", "Bakery 是「节假日 = 信号」的教科书案例")
pic(s, FIG / "fig11_hierarchy_mape.png", Inches(0.55), Inches(1.5), w=Inches(7.6))
rect(s, Inches(8.4), Inches(1.55), Inches(4.35), Inches(4.9), LIGHT)
text(s, Inches(8.65), Inches(1.8), Inches(3.9), Inches(0.6), "Bakery 分品类", size=17, color=NAVY, bold=True)
lines(s, Inches(8.65), Inches(2.45), Inches(3.9), Inches(3.0), [
    ("朴素 / SARIMA", "  76.3% / 78.3%"),
    ("SARIMAX", "  30.4%"),
])
text(s, Inches(8.65), Inches(4.4), Inches(3.9), Inches(1.6),
     "烘焙需求由圣诞/感恩节峰驱动，纯周模型无法触及，显式因子精准捕捉。", size=12, color=SUB)

# 7 — Takeaway
s = slide()
title_bar(s, "结论与启示", "可信数据，不仅是正确 —— 而是启发")
bullets(s, Inches(0.55), Inches(1.5), Inches(8.0), Inches(5.2), [
    ("节假日是最强可预测信号", " —— 优先于价格促销做库存/排班"),
    ("品类分层建模", " —— Bakery 季节尖峰型 vs Meat 刚性型"),
    ("门店节奏一致、体量不同", " —— 总盘按体量分摊 + 门店级微调"),
    ("每个数字可追溯", " —— reports/*.json + 确定性脚本一键重放"),
], size=16)
rect(s, Inches(8.8), Inches(1.6), Inches(4.0), Inches(4.2), LIGHT)
lines(s, Inches(9.0), Inches(1.85), Inches(3.6), Inches(3.6), [
    ("SARIMA", "  15.1%"),
    ("SARIMAX", "  10.5%"),
    ("RMSE", "  1743 → 1286"),
    ("AIC", "  同步下降"),
])

prs.save(str(OUT))
print("PPTX written to", OUT)